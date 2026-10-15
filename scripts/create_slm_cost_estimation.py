# -*- coding: utf-8 -*-
"""
SLM 개발대가 산정 생성기 — 단일 FP유형(간이법) + 엔지니어링 내역
- 양식: SW사업 대가산정 가이드(KOSA) 단일 FP유형 예제 + 엔지니어링 내역(R5) 예제
- 기능 목록 기준일: 2026-10-15 (메뉴 57종·채팅 인텐트 81종·Migration 0143 기준,
  폐기 기능(상황보고·현장 모드·현장 지식·인수인계) 제외)
실행: scratchpad venv 등 openpyxl 있는 파이썬으로
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = sys.argv[1] if len(sys.argv) > 1 else "SLM_개발대가산정_단일FP_엔지니어링.xlsx"

wb = Workbook()

# ── 스타일 ──
HDR_FONT = Font(name="맑은 고딕", bold=True, size=10, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="2F5496")
CAT_FILL = PatternFill("solid", fgColor="D6E4F0")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
SUB_FILL = PatternFill("solid", fgColor="E2EFDA")
TITLE_FONT = Font(name="맑은 고딕", bold=True, size=13)
NF = Font(name="맑은 고딕", size=10)
BF = Font(name="맑은 고딕", bold=True, size=10)
SMALL = Font(name="맑은 고딕", size=9, color="666666")
thin = Side(style="thin", color="B4C6E7")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
L = Alignment(horizontal="left", vertical="center", wrap_text=True)
R = Alignment(horizontal="right", vertical="center")


def hdr(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cl = ws.cell(row=row, column=c)
        cl.font = HDR_FONT; cl.fill = HDR_FILL; cl.alignment = C; cl.border = BD


def box(ws, row, c1, c2, fill=None, font=None):
    for c in range(c1, c2 + 1):
        cl = ws.cell(row=row, column=c)
        cl.font = font or NF; cl.border = BD
        if fill: cl.fill = fill


def widths(ws, ww):
    for c, w in enumerate(ww, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


FP_W = {"ILF": 7.5, "EIF": 5.4, "EI": 4.0, "EO": 5.2, "EQ": 3.9}
FP_UNIT_PRICE = 605784  # 기능점수당 단가 (KOSA 2024 가이드. 최신 공표 시 갱신)

# ═════════════════════════════════════════════════════════════════
# 기능 목록 — (세부업무, 단위프로세스, 설명, FP유형, 사용자관점)
# ═════════════════════════════════════════════════════════════════
F = []
def sec(name): F.append(("SEC", name, "", "", ""))
def p(proc, desc, tp, user=""):  F.append(("", proc, desc, tp, user))

sec("1. 실시간 감시")
p("통합 대시보드", "시설유형별 KPI·이상 시설·유량 불균형·설비 장애·최근 경보 종합 표출", "EO", "Y")
p("대시보드 집계 데이터", "이상 스캔·물수지·장애 집계 캐시 (디스크 영속화 포함)", "ILF", "Y")
p("시작 페이지", "로그인 후 첫 화면(랜딩) 구성·바로가기", "EQ", "")
p("GIS 관망도", "오프라인 지도 위 시설·관로·블록 레이어와 실시간 상태 표출", "EO", "Y")
p("GIS 시설 상세 팝업", "시설 클릭 → 현재값·경보·바로가기 팝업", "EQ", "")
p("GIS 타임라인 스크러버", "24h 되감기/재생·배속·알람 도트", "EO", "")
p("GIS 시설 좌표 관리", "시설 좌표 등록·수정 (지도 설정)", "EI", "")
p("실시간 계통도", "SVG 계통도·유량 비례 엣지·흐름 애니메이션·미니맵·검색", "EO", "Y")
p("계통도 노드 상세", "노드별 Q/H/P 수치 오버레이·트렌드 패널", "EQ", "")
p("태그 모니터링", "모니터링 그룹 현재값·이상 카테고리 9종 필터·정렬·우클릭 트렌드", "EO", "Y")
p("시설별 모니터링(배수지)", "배수지 트렌드·HH/LL 가이드·재생·시간 네비게이션", "EO", "Y")
p("시설별 모니터링(가압장)", "가압장 펌프·압력·유량 트렌드 모니터링", "EO", "")
p("시설별 모니터링(블록)", "블록 유량·압력 트렌드 모니터링", "EO", "")
p("모니터링 임계 설정", "HH/LL 임계 듀얼모드(상수/태그참조) 설정", "EI", "")
p("상황실 키오스크", "기존 화면 iframe 상주 순환·크롬 제거·하단 바(시계·티커)", "EO", "Y")
p("키오스크 경보 인터럽트", "새 경고 감지 시 경보 화면 전환·배너·재인터럽트 방지", "EO", "")
p("키오스크 설정", "순환 화면 목록·순서, 야간 밝기, KPI 스트립 설정", "EI", "")

sec("2. 경보 관리")
p("경보 데이터", "경보 발생·해제·확인·판정 이력 (하이퍼테이블 연계)", "ILF", "Y")
p("진행중 경보 현황", "진행중 경보 목록·반복 경보 접기(반복 N회)·심각도 구분", "EO", "Y")
p("경보 이력 조회", "기간·등급·시설 필터 경보 이력 조회", "EQ", "")
p("반복 경보 순위", "채터링/만성 판정·순위표·조치 구분 제시", "EO", "")
p("경보 확인 처리", "확인/해제 기록·최초 확인자 책임 추적", "EI", "")
p("미확인 경보 상신", "경고 N분 미확인 시 전체 채널 1회 멱등 상신", "EO", "")
p("경보 분석 상세", "근본원인 분류·위기대응 섹션·연관 비상연락처 매칭", "EO", "Y")
p("근본원인 자동 분류", "경보 발생 원인 자동 분류 (주기 작업)", "EO", "")
p("오탐 라벨 판정", "3분법(real/false/check) 1클릭 판정·백필 체계", "EI", "")
p("오탐 분석 통계", "주별 오탐률 추이·분류별 현황", "EO", "")
p("임계 도달 예측", "수위 LEI↔LEC 페어 선형 외삽·ETA 구간 표시", "EO", "")
p("알람 캘린더", "일자별 경보 발생 캘린더 조회", "EQ", "")
p("알람 팝업(위기대응 모달)", "새 경고 감지 → 모달 강제(주의는 toast)·4액션", "EO", "Y")
p("누수 의심 알림", "야간최소유량 CUSUM 자동 감지·선정 사유 서술·확인 처리", "EO", "Y")
p("누수 알림 확인 처리", "누수 의심 건 확인·상태 변경", "EI", "")
p("작업관리", "작업 등록·조회·알람 억제 시간 설정", "EI", "Y")
p("작업 이력 데이터", "작업·조치 이력 (장애 기록과 연계)", "ILF", "")
p("펌프 가동 변경", "펌프 제어모드·가동 변경 기록 관리", "EI", "")
p("시나리오 비교", "경보 시나리오 간 비교 분석 화면", "EO", "")

sec("3. AI 채팅 어시스턴트")
p("채팅 질의 처리(SSE)", "3단계 인텐트 분류(키워드→벡터→SLM)·SQL 실행·SSE 스트리밍, 81 인텐트", "EO", "Y")
p("인텐트 정의 데이터", "인텐트 레지스트리(JSON 선언·분류 키워드·동적 SQL)", "ILF", "Y")
p("인텐트 벡터 검색", "로컬 임베딩 모델 기반 유사도 분류 (외부 API 없음)", "EIF", "")
p("채팅 시각화 렌더링", "그래프·테이블·다이어그램·문서 카드 자동 라우팅", "EO", "Y")
p("채팅 대화 이력(서버)", "계정별 대화 그룹·메시지 서버 저장·본문 검색·30일 purge", "ILF", "Y")
p("채팅 이력 조회", "대화 목록·검색·차트 재렌더링", "EQ", "")
p("추천 질의", "인텐트별 후속 추천·현장명 자동 보강", "EO", "")
p("오답 피드백 루프", "원하는 답이 아닌가요 → 후보 제시·피드백 수집·보정", "EI", "Y")
p("멀티턴 후속질문", "짧은 후속질문 인텐트 상속·지표 명사 가드", "EO", "")
p("멀티모달 사진 진단", "채팅 사진 업로드 → 장비 식별·매뉴얼 RAG·케이스 DB 통합 진단", "EO", "Y")
p("고장 기록 초안(채팅)", "자연어 고장 등록 → 시설·설비 자동 매칭·확인 카드", "EI", "Y")
p("조치 완료 기록(채팅)", "진행중 장애 자연어 조치 완료 처리", "EI", "")
p("점검 기록(채팅)", "일상/요청 점검 기록·유발 고장 링크", "EI", "")
p("보고서 초안(채팅)", "기간·유형 자동 추출 → 보고서 자동 초안 생성", "EO", "")
p("진단 근거 팩", "룰 라우터가 조회 도구 3종 병렬 실행 → 근거 표시·LLM 참고 의견", "EO", "")
p("음성 입력", "로컬 Whisper STT·도메인 용어 바이어스", "EI", "")
p("비상연락처 질의", "채팅에서 비상연락처 조회 인텐트", "EQ", "")
p("FAQ 동적 제안", "FAQ 기반 예시 질문 동적 노출", "EQ", "")

sec("4. 트렌드 분석")
p("트렌드 카탈로그", "트렌드 정의 단일 정본(시드 보호·모니터링 큐레이션)", "ILF", "Y")
p("사용자 트렌드 생성", "사용자 정의 트렌드 등록·수정 (채팅 즉시 조회 연동)", "EI", "Y")
p("트렌드 차트 조회", "기간·이동평균·재생·확장 보기", "EO", "Y")
p("평소 대비 비교", "GBT 기대값 기반 z-score 판정·KPI 배지·수위 하한 판정", "EO", "Y")
p("향후 전망", "Chronos-Bolt 예측·밴드 렌더·시간대 중앙값 프로파일 폴백", "EO", "")
p("기준선 GBT 학습", "주 1회 재훈련 CLI·성능 평가 지표 적재", "EO", "")
p("IForest 이상탐지", "디스크 영속화 모델·주 1회 재학습·평가 지표", "EO", "")
p("야간최소유량 집계", "일 단위 NMF 사전집계·self-healing 백필", "EO", "")
p("태그 원시데이터", "TimescaleDB 하이퍼테이블(압축 정책·연속 집계)", "ILF", "Y")

sec("5. 장애·설비 관리")
p("설비 정보", "설비 마스터·제원(meta)·내용연수", "ILF", "Y")
p("고장 케이스 DB", "고장 사례·조치 지식 데이터", "ILF", "Y")
p("고장 케이스 관리", "케이스 등록·수정·xlsx 일괄 업로드", "EI", "")
p("장애 이력 조회", "설비별 장애·조치 이력·통계", "EQ", "Y")
p("설비 건강성 대시보드", "알람↔조치 대조 분석·오탐/미검지·교체 후보", "EO", "Y")
p("교체 우선순위 Top5", "내용연수·MTBF·재발 3신호 융합 카드", "EO", "")
p("설비 신뢰성(MTBF)", "설비 유형별 MTBF·고장 간격 분석", "EO", "")
p("점검 주기 도래", "유형 단위 점검 주기 마스터·도래 계산", "EQ", "")
p("설비 매뉴얼 관리", "매뉴얼 PDF 업로드·인덱싱·핫리로드", "EI", "")
p("매뉴얼 RAG 검색", "장·절 청크 임베딩 검색 (로컬)", "EIF", "")
p("명판 비전 스캔", "명판 촬영 → 동적 키-값 추출·표준 키 승격·제원 반영", "EI", "")

sec("6. 보고서·업무")
p("장애 조치 보고서", "항목 편집·AI 요약·확정·인쇄(공문 양식)", "EI", "Y")
p("일 점검 보고서", "점검 기록 집계·고장보고 포함·확정·인쇄", "EI", "Y")
p("보고서 데이터", "보고서·항목·확정 상태 저장", "ILF", "")
p("보고서 조회·출력", "보고서 목록·상세·새 창 인쇄(PDF)", "EO", "")
p("업무 메모", "메모 등록·검색·달력 보기 (본인+마스터 삭제)", "EI", "")
p("일정 알림", "달력 등록 → 30s 폴링 전역 팝업·확인 전 재표시", "EO", "")

sec("7. 관망 수리해석(EPANET)")
p("EPANET 모델 데이터", "INP 네트워크 모델·시뮬 결과 시계열", "ILF", "")
p("실측 유량 주입", "시설별 outflow 매핑 → demand IDW 보간 주입", "EO", "")
p("수리 시뮬레이션 실행", "EPANET 시뮬 실행·결과 저장 (자동 주기 포함)", "EO", "Y")
p("시뮬 vs 실측 편차 분석", "시설 단위 편차 패널·GIS 오버레이", "EO", "")
p("상류 영향 분석", "상류 시설 장애 영향 추정·채팅 연동", "EO", "")

sec("8. 통신망 관리")
p("통신 장비 정보", "네트워크 장비 마스터(IP·프로토콜·역할)", "ILF", "Y")
p("통신 링크 정보", "장비 간 연결·프로토콜 정의", "ILF", "")
p("네트워크 토폴로지", "계층형 토폴로지·장비 상태·RTT 감시", "EO", "Y")
p("통신 이상 감지", "무응답·통신 이상 자동 감지·경보 연계", "EO", "")
p("장비 CRUD", "장비 등록·수정·삭제 (표+캔버스 동일 정본)", "EI", "")
p("링크 에디터", "드래그 연결 등록·더블클릭 수정·배치 저장", "EI", "Y")
p("링크 에디터 lint", "미연결·IP중복·겹침·동기이탈 검사 배지", "EQ", "")
p("배치 정렬", "계층 정렬·토폴로지 정렬(모니터링 동일 작화)·정렬 5종", "EO", "")

sec("9. 메신저·협업")
p("메시지 데이터", "1:1/그룹 대화·읽음·첨부 메타", "ILF", "Y")
p("운영자 채팅", "1:1/그룹 채팅·도착 알림(10s 폴링)", "EI", "Y")
p("파일 전송", "사진·영상·문서 첨부 전송", "EI", "")
p("음성 통화", "WebRTC P2P 음성 통화·REST 폴링 시그널링", "EO", "Y")
p("영상 통화", "WebRTC 영상 통화·화면 구성", "EO", "")
p("회의 통화", "3~4인 풀메시 회의·초대 벨·수락 합류", "EO", "")
p("외부망 TURN 옵션", "coturn 릴레이 자격 발급·이중 토글", "EO", "")

sec("10. 구축(셋업)")
p("구축 워크스페이스", "트리+레이어 토글(용수/통신)+공용 인스펙터 통합 진입점", "EQ", "Y")
p("시설 등록 마법사", "4단계 시설 등록(기초정보→태그→배치)", "EI", "Y")
p("시설 기초정보", "배수지·가압장·감압·블록 4종 기초정보", "ILF", "Y")
p("시설 기초정보 CRUD", "시설 등록·수정·삭제·사진/매뉴얼 연결", "EI", "")
p("태그 설정", "태그 마스터 관리·일괄 업로드·온보딩", "EI", "Y")
p("DATAINFO 변환룰", "룰 4계층(정규식/사전/문맥/확정)·미리보기·선별 적용·롤백", "EI", "")
p("변환룰 데이터", "룰 정의·적용 이력", "ILF", "")
p("장비 설정", "장비 등록·태그 연결·사진 등록", "EI", "")
p("트렌드 설정", "트렌드 카탈로그 관리(시드 보호)", "EI", "")
p("계통도 캔버스", "노드 드래그·연결·undo/redo·내보내기(PNG/SVG)", "EI", "Y")
p("계통도 좌표 데이터", "노드 배치 정본(경위도·px 변환)", "ILF", "")
p("자동 배치·정합 lint", "relayout(신규/전체)·누락/고아/순환 검사", "EO", "")
p("GIS 설정", "지도 중심·줌·좌표 관리", "EI", "")
p("비상연락처 설정", "카테고리별 연락처 관리", "EI", "")
p("비상연락처 데이터", "연락처·카테고리 매핑", "ILF", "")
p("시설 약칭 설정", "채팅 매칭용 시설 별칭 관리", "EI", "")
p("구축 완결성 검수", "검사 6종(기초정보/임계/datainfo/계통도/품질/EPANET) 집계", "EO", "Y")

sec("11. 관리·공통")
p("로그인·인증", "JWT 발급·세션 가드·bcrypt 점진 마이그레이션", "EI", "Y")
p("사용자 데이터", "사용자·권한 정보", "ILF", "")
p("사용자 관리", "사용자 등록·수정·권한 부여", "EI", "")
p("메뉴 관리", "메뉴 트리·권한별 접근 제어", "EI", "")
p("메뉴 데이터", "메뉴 트리·권한 매핑", "ILF", "")
p("사이트 설정", "기능 토글(SKU)·지도·AI 파라미터·키오스크 설정 통합", "EI", "Y")
p("설정 데이터", "사이트 설정 코드(tb_comm_code)", "ILF", "")
p("FAQ 관리", "FAQ 등록·노출 관리", "EI", "")
p("시설 파일 관리", "시설 사진·계통도·매뉴얼 파일 연결 관리", "EI", "")
p("파일 저장 데이터", "파일 메타·저장소 연동", "ILF", "")
p("채팅 피드백 관리", "오답 피드백 목록·처리 현황", "EQ", "")
p("LLM 서술 관찰", "LLM 참고 의견 생성 이력 관찰", "EQ", "")
p("보고서 카테고리 관리", "보고서 분류 체계 관리", "EI", "")
p("AI 모델 평가", "GBT·IForest 평가 지표 통합 화면", "EQ", "")
p("감사 로그", "주요 행위 감사 추적 조회", "EQ", "")
p("감사 로그 데이터", "감사 이벤트 저장", "ILF", "")
p("채팅 카드 갤러리", "카드 유형별 미리보기 갤러리", "EQ", "")
p("시스템 버전·모듈", "모듈 버전·live health·라이선스·OSS 고지·업데이트 번들", "EO", "Y")
p("모듈 버전 데이터", "모듈·라이선스·업데이트 이력", "ILF", "")
p("테마·레이아웃 설정", "브랜드 컬러·사이드바/탑바·다크/라이트", "EI", "")

sec("12. 데이터 수집·품질")
p("SCADA 수집 연계", "Node-RED 경유 계측 데이터 수집 연계", "EIF", "Y")
p("알람 조건 엔진 연계", "수위/압력/펌프 조건검토·경보 발생/해제 연계", "EIF", "Y")
p("태그 품질 계층", "포화·고착·무신호·DI반전·값이탈 5종 검사·판정 게이트", "EO", "Y")
p("품질 데이터", "태그 품질 상태 저장", "ILF", "")
p("사전집계 관리", "1h/5분 연속 집계·압축 정책 운영", "EO", "")

# ═════════════════════════════════════════════════════════════════
# Sheet 1: FP산정(간이법)
# ═════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "FP산정(간이법)"
ws1.sheet_properties.tabColor = "2F5496"
ws1.cell(row=1, column=2, value="기능점수 산정 (간이법 평균복잡도 적용)").font = TITLE_FONT

r = 3
heads = ["", "기능명", "", "", "데이터 및 트랜잭션 기능", "FP 산출", "사용자관점 기능", "비고"]
for c, h in enumerate(heads, 1):
    ws1.cell(row=r, column=c, value=h)
hdr(ws1, r, 2, 8)
r = 4
for c, h in enumerate(["", "①어플리케이션명", "②세부 업무명", "③단위프로세스명", "단위프로세스 설명", "④FP유형", "⑤가중치", ""], 1):
    ws1.cell(row=r, column=c, value=h)
hdr(ws1, r, 2, 8)

W_FORMULA = ('=IF(F{r}="","",IF(F{r}="EI",4,IF(F{r}="EO",5.2,IF(F{r}="EQ",3.9,'
             'IF(F{r}="ILF",7.5,IF(F{r}="EIF",5.4,""))))))')
row = 5
first_data_row = row
app_written = False
for kind, a, b, c_, d in F:
    if kind == "SEC":
        ws1.cell(row=row, column=3, value=a).font = BF
        if not app_written:
            ws1.cell(row=row, column=2, value=" < SLM 스마트 물관리 시스템 > ").font = BF
            app_written = True
        box(ws1, row, 2, 8, fill=CAT_FILL)
        ws1.cell(row=row, column=3).alignment = L
        sec_row = row
        row += 1
        continue
    ws1.cell(row=row, column=4, value=a)
    ws1.cell(row=row, column=5, value=b)
    ws1.cell(row=row, column=6, value=c_)
    ws1.cell(row=row, column=7, value=W_FORMULA.format(r=row))
    ws1.cell(row=row, column=8, value=d)
    box(ws1, row, 2, 8)
    ws1.cell(row=row, column=4).alignment = L
    ws1.cell(row=row, column=5).alignment = L
    row += 1
last_data_row = row - 1
ws1.cell(row=row, column=2, value="합계").font = BF
ws1.cell(row=row, column=7,
         value=f"=SUM(G{first_data_row}:G{last_data_row})").font = BF
box(ws1, row, 2, 8, fill=TOTAL_FILL, font=BF)
FP_TOTAL = f"'FP산정(간이법)'!G{row}"
widths(ws1, [2, 22, 18, 24, 58, 9, 9, 11, 10])

# ═════════════════════════════════════════════════════════════════
# Sheet 2: SW개발비 산정
# ═════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("SW개발비 산정")
ws2.sheet_properties.tabColor = "C00000"
ws2.cell(row=1, column=2, value='"SLM 스마트 물관리 시스템" 소프트웨어 개발비 산정').font = TITLE_FONT
ws2.cell(row=2, column=9, value="(소프트웨어사업 대가산정 가이드, KOSA)").font = SMALL

ws2.cell(row=4, column=2, value="○ 개발원가 산정").font = BF
ws2.cell(row=4, column=10, value="(단위: 원)").font = SMALL
for c, h in enumerate(["", "총기능점수", "기능점수당 단가", "보정계수", "", "", "", "", "개발원가"], 1):
    ws2.cell(row=5, column=c, value=h)
hdr(ws2, 5, 2, 9)
for c, h in enumerate(["", "", "", "규모", "연계복잡성", "성능", "운영환경", "보안성", ""], 1):
    ws2.cell(row=6, column=c, value=h)
hdr(ws2, 6, 2, 9)
ws2.cell(row=7, column=2, value=f"={FP_TOTAL}")
ws2.cell(row=7, column=3, value=FP_UNIT_PRICE)
ws2.cell(row=7, column=4, value="=J21")
ws2.cell(row=7, column=5, value="=J22")
ws2.cell(row=7, column=6, value="=J23")
ws2.cell(row=7, column=7, value="=J24")
ws2.cell(row=7, column=8, value="=J25")
ws2.cell(row=7, column=9, value="=ROUND(PRODUCT(B7:H7),0)")
box(ws2, 7, 2, 9)
ws2.cell(row=8, column=2, value="합계(보정 후 개발원가)").font = BF
ws2.cell(row=8, column=9, value="=I7").font = BF
box(ws2, 8, 2, 9, fill=TOTAL_FILL, font=BF)
ws2.cell(row=9, column=2, value="이윤 (개발원가의 25% 이내)")
ws2.cell(row=9, column=8, value=0.2)
ws2.cell(row=9, column=9, value="=ROUND(I8*H9,0)")
box(ws2, 9, 2, 9)
ws2.cell(row=10, column=2, value="직접경비")
ws2.cell(row=10, column=9, value="=I17")
box(ws2, 10, 2, 9)
ws2.cell(row=11, column=2, value="소프트웨어 개발비 (부가세 별도)").font = BF
ws2.cell(row=11, column=9, value="=SUM(I8:I10)").font = BF
box(ws2, 11, 2, 9, fill=SUB_FILL, font=BF)

ws2.cell(row=13, column=2, value="○ 직접경비").font = BF
for c, h in enumerate(["", "구분", "산출내역", "", "", "", "", "", "금액"], 1):
    ws2.cell(row=14, column=c, value=h)
hdr(ws2, 14, 2, 9)
ws2.cell(row=15, column=3, value="직접경비 요소 없음 (폐쇄망 온프레미스 — 발주자 제공 인프라 기준)")
ws2.cell(row=15, column=9, value=0)
box(ws2, 15, 2, 9)
ws2.cell(row=17, column=2, value="합 계").font = BF
ws2.cell(row=17, column=9, value="=SUM(I15:I16)").font = BF
box(ws2, 17, 2, 9, fill=TOTAL_FILL, font=BF)

ws2.cell(row=19, column=2, value="○ 보정계수 산정").font = BF
for c, h in enumerate(["", "보정계수 구분", "복잡도 및 난이도 수준", "", "", "", "", "", "", "보정계수"], 1):
    ws2.cell(row=20, column=c, value=h)
hdr(ws2, 20, 2, 10)
coef_rows = [
    ("SW규모", "= 0.4057 × (ln(FP) − 7.1978)² + 0.8878  (500FP 미만 1.28, 3,000FP 초과 1.153)",
     "=IF(B7<500,1.28,IF(B7>3000,1.153,ROUND(0.4057*(LN(B7)-7.1978)^2+0.8878,4)))"),
    ("연계복잡성", "1. 타 기관 연계 없음 (폐쇄망 단독 운영 — SCADA 연계는 내부 연계)", 0.88),
    ("성능 요구수준", "3. 응답시간·처리율이 피크타임에 중요하며 처리 시한이 명시 (실시간 감시·알람)", 1.0),
    ("운영환경 호환성", "3. 호환성 요구 있음 — 유사 HW·SW 환경(고객별 온프레미스 설치)", 1.06),
    ("보안성 요구수준", "2. 2가지 요구사항 포함 (암호화(JWT·bcrypt)·시큐어코딩)", 1.0),
]
for i, (name, desc, val) in enumerate(coef_rows):
    rr = 21 + i
    ws2.cell(row=rr, column=2, value=name)
    ws2.cell(row=rr, column=3, value=desc)
    ws2.cell(row=rr, column=10, value=val)
    box(ws2, rr, 2, 10)
    ws2.cell(row=rr, column=3).alignment = L

ws2.cell(row=27, column=2, value="○ FP 집계").font = BF
for c, h in enumerate(["", "구분", "ILF", "EIF", "EI", "EO", "EQ", "계"], 1):
    ws2.cell(row=28, column=c, value=h)
hdr(ws2, 28, 2, 8)
types = ["ILF", "EIF", "EI", "EO", "EQ"]
ws2.cell(row=29, column=2, value="복잡도(가중치)")
for i, t in enumerate(types):
    ws2.cell(row=29, column=3 + i, value=FP_W[t])
box(ws2, 29, 2, 8)
ws2.cell(row=30, column=2, value="기능수")
for i, t in enumerate(types):
    ws2.cell(row=30, column=3 + i,
             value=f"=COUNTIF('FP산정(간이법)'!$F${first_data_row}:$F${last_data_row},\"{t}\")")
ws2.cell(row=30, column=8, value="=SUM(C30:G30)")
box(ws2, 30, 2, 8)
ws2.cell(row=31, column=2, value="기능점수")
for i, t in enumerate(types):
    ws2.cell(row=31, column=3 + i,
             value=f"=SUMIF('FP산정(간이법)'!$F${first_data_row}:$F${last_data_row},\"{t}\","
                   f"'FP산정(간이법)'!$G${first_data_row}:$G${last_data_row})")
ws2.cell(row=31, column=8, value="=SUM(C31:G31)")
box(ws2, 31, 2, 8, font=BF)
widths(ws2, [2, 26, 16, 12, 12, 12, 12, 12, 16, 10])

# ═════════════════════════════════════════════════════════════════
# Sheet 3: (참조)보정계수
# ═════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("(참조)보정계수")
ws3.cell(row=1, column=2, value="보정계수 참조표 — FP방식(간이법)").font = TITLE_FONT
ref = [
    ("연계복잡성", [("1. 타 기관 연계 없음", 0.88), ("2. 1~2개의 타 기관 연계", 0.94),
     ("3. 3~5개의 타 기관 연계", 1.0), ("4. 6~10개의 타 기관 연계", 1.06),
     ("5. 10개 초과의 타 기관 연계", 1.12)]),
    ("성능 요구수준", [("1. 응답성능에 대한 특별한 요구사항이 없다.", 0.91),
     ("2. 요구사항이 있으나 특별한 조치는 불필요", 0.95),
     ("3. 응답시간·처리율이 피크타임에 중요, 처리 시한 명시", 1.0),
     ("4. 모든 업무시간에 중요, 처리 시한 명시", 1.05),
     ("5. 엄격 — 설계·구현 단계 성능 분석도구 필요", 1.09)]),
    ("운영환경 호환성", [("1. 요구사항 없음", 0.94), ("2. 동일 HW·SW 환경", 1.0),
     ("3. 유사 HW·SW 환경", 1.06), ("4. 이질적 HW·SW 환경", 1.13),
     ("5. 4 + 운영 절차 문서화·사전 모의훈련 요구", 1.19)]),
    ("보안성 수준", [("1. 1가지 보안 요구사항", 0.97), ("2. 2가지", 1.0),
     ("3. 3가지", 1.03), ("4. 4가지", 1.06), ("5. 5가지 이상", 1.08)]),
]
ws3.cell(row=3, column=2, value="○ 규모 보정계수").font = BF
ws3.cell(row=4, column=2, value="SW규모")
ws3.cell(row=4, column=3,
         value="= 0.4057 × (ln(기능점수) − 7.1978)² + 0.8878  (단, 500FP 미만 1.28 / 3,000FP 초과 1.153)")
box(ws3, 4, 2, 8)
rr = 6
ws3.cell(row=rr, column=2, value="○ 애플리케이션 복잡도 보정계수").font = BF
rr += 1
for group, items in ref:
    for i, (desc, v) in enumerate(items):
        if i == 0:
            ws3.cell(row=rr, column=2, value=group).font = BF
        ws3.cell(row=rr, column=3, value=desc)
        ws3.cell(row=rr, column=8, value=v)
        box(ws3, rr, 2, 8)
        ws3.cell(row=rr, column=3).alignment = L
        rr += 1
widths(ws3, [2, 18, 70, 8, 8, 8, 8, 10])

# ═════════════════════════════════════════════════════════════════
# Sheet 4: 직접경비 (참조)
# ═════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("직접경비")
items4 = [
    ("컴퓨터시스템 사용료", "발주자가 제공하지 않아 자체/제3자 조달 시"),
    ("소프트웨어 도구 사용료", "발주자 요구에 의한 SW 패키지 사용료"),
    ("지급이자", "선투자 후정산 사업의 경우"),
    ("특정기술 전문가 비용", "발주자 요구의 특정기술 도입 관련"),
    ("여비", "운임·현지교통비·숙박비 등"),
    ("특수자료비", "특허·노하우 사용료"),
    ("인쇄·청사진비", "산출물·보고서 인쇄·복사비"),
    ("자료조사비", "문헌·전문도서 구입"),
    ("기자재시험비", "기자재 시험 소요 비용"),
    ("위탁비(외주비)", "일부 외부 위탁 비용"),
    ("현장운영비", "보조요원 급여·현장사무실 임차·운영비"),
    ("모형제작비", "프로토타입 제작 비용"),
    ("기타", "그 밖에 특별히 소요되는 직접비용"),
]
ws4.cell(row=1, column=1, value="직접경비 항목 (참조)").font = TITLE_FONT
for c, h in enumerate(["항목", "설명"], 1):
    ws4.cell(row=2, column=c, value=h)
hdr(ws4, 2, 1, 2)
for i, (a, b) in enumerate(items4):
    ws4.cell(row=3 + i, column=1, value=a)
    ws4.cell(row=3 + i, column=2, value=b)
    box(ws4, 3 + i, 1, 2)
    ws4.cell(row=3 + i, column=2).alignment = L
widths(ws4, [34, 80])

# ═════════════════════════════════════════════════════════════════
# Sheet 5: 엔지니어링 내역 (설치·구축 — SW개발비 외 작업)
# ═════════════════════════════════════════════════════════════════
ws5 = wb.create_sheet("엔지니어링 내역")
ws5.sheet_properties.tabColor = "548235"
ws5.cell(row=1, column=1, value="■ SLM 구축·설치 엔지니어링 내역 (설계금액)").font = TITLE_FONT

# 노임 기준 (SW기술자 평균임금 공표 2025, 월 20.5일 · 1일=0.048 M/M)
WAGE_ROW_IT, WAGE_ROW_DEV, WAGE_ROW_OPS = 40, 41, 42
for c, h in enumerate(["구분", "항목", "작업 내용", "M/M", "", "", "금액(원)", "", "", "기간", "비고"], 1):
    ws5.cell(row=2, column=c, value=h)
hdr(ws5, 2, 1, 11)
for c, h in enumerate(["", "", "", "IT지원", "응용SW개발", "정보시스템운용", "IT지원", "응용SW개발", "정보시스템운용", "", ""], 1):
    ws5.cell(row=3, column=c, value=h)
hdr(ws5, 3, 1, 11)

# (구분, 항목, 내용, it_days, dev_days, ops_days, 기간, 비고)
ENG = [
    ("설치", "서버 스택 설치·구성", "Docker 스택(DB·백엔드·프런트·Node-RED)·프로덕션 빌드·TLS 구성", 0, 2, 5, "1주", ""),
    ("", "DB 구축", "스키마·마이그레이션 적용·기초 시드·압축 정책", 0, 2, 3, "1주", ""),
    ("", "AI 모델 설치", "Ollama 모델(비전·임베딩) 오프라인 설치·검증", 0, 0, 2, "2일", "폐쇄망 반입"),
    ("", "모델 웨이트 번들 반입", "Chronos·Whisper 번들 설치·sha256 검증", 0, 0, 2, "2일", ""),
    ("", "GIS 지도 번들 구축", "관할 영역 pmtiles 추출·글리프 반입·검증", 0, 3, 0, "3일", "관할별 재추출"),
    ("데이터 구축", "태그 마스터 구축", "태그 등록·DATAINFO 변환룰 적용·검수 (약 2,700태그)", 10, 10, 0, "2주", "룰 재현율 채점"),
    ("", "시설 기초정보 구축", "배수지·가압장·감압·블록 기초정보 (약 120개소)", 10, 0, 0, "2주", ""),
    ("", "계통도 구축", "수계 관계(flow-map)·캔버스 배치·정합 lint 0건", 0, 5, 0, "1주", ""),
    ("", "GIS 좌표 구축", "시설 좌표 등록·현장 검증", 5, 0, 0, "1주", ""),
    ("", "통신망 구축", "장비·링크 등록 (약 300대)·토폴로지 검증", 7, 0, 0, "1.5주", ""),
    ("", "EPANET 모델 구축", "INP 모델 작성·실측 유량 매핑·시뮬 검증", 0, 10, 0, "2주", "B1 SKU 시"),
    ("연계", "SCADA 수집 연계", "Node-RED 수집 플로우 구성·태그 매핑", 0, 10, 5, "2주", ""),
    ("", "알람 조건 엔진 구성", "수위/압력/펌프 조건검토 플로우 현장 임계 보정", 0, 10, 0, "2주", ""),
    ("시험·인수", "구축 완결성 검수", "검수 6종 warn 0 달성·품질 게이트 통과", 0, 5, 0, "1주", "인수 기준"),
    ("", "통합 시험", "채팅 스모크·알람 시나리오·통화(LTE) 시험", 0, 5, 0, "1주", ""),
    ("", "교육·인수인계", "운영자 교육·운영 문서 인수", 5, 0, 0, "1주", ""),
]
DAY = 0.048
row = 4
first_eng = row
for gu, item, desc, it_d, dev_d, ops_d, dur, note in ENG:
    ws5.cell(row=row, column=1, value=gu)
    ws5.cell(row=row, column=2, value=item)
    ws5.cell(row=row, column=3, value=desc)
    if it_d:
        ws5.cell(row=row, column=4, value=f"={DAY}*{it_d}")
        ws5.cell(row=row, column=7, value=f"=D{row}*($B${WAGE_ROW_IT}*20.5)")
    if dev_d:
        ws5.cell(row=row, column=5, value=f"={DAY}*{dev_d}")
        ws5.cell(row=row, column=8, value=f"=E{row}*($B${WAGE_ROW_DEV}*20.5)")
    if ops_d:
        ws5.cell(row=row, column=6, value=f"={DAY}*{ops_d}")
        ws5.cell(row=row, column=9, value=f"=F{row}*($B${WAGE_ROW_OPS}*20.5)")
    ws5.cell(row=row, column=10, value=dur)
    ws5.cell(row=row, column=11, value=note)
    box(ws5, row, 1, 11)
    ws5.cell(row=row, column=2).alignment = L
    ws5.cell(row=row, column=3).alignment = L
    row += 1
last_eng = row - 1
ws5.cell(row=row, column=2, value="합계").font = BF
for col in ["D", "E", "F", "G", "H", "I"]:
    ws5.cell(row=row, column="DEFGHI".index(col) + 4,
             value=f"=SUM({col}{first_eng}:{col}{last_eng})")
box(ws5, row, 1, 11, fill=TOTAL_FILL, font=BF)
sum_row = row
row += 2
ws5.cell(row=row, column=1, value="직접인건비").font = BF
ws5.cell(row=row, column=3, value="설치·데이터 구축·연계·시험 엔지니어링 합계")
ws5.cell(row=row, column=8, value=f"=G{sum_row}+H{sum_row}+I{sum_row}").font = BF
box(ws5, row, 1, 11)
di_row = row
row += 1
ws5.cell(row=row, column=1, value="제경비")
ws5.cell(row=row, column=3, value="직접인건비의 110~120% (상한 120% 적용)")
ws5.cell(row=row, column=8, value=f"=ROUND(H{di_row}*1.2,0)")
box(ws5, row, 1, 11)
je_row = row
row += 1
ws5.cell(row=row, column=1, value="기술료")
ws5.cell(row=row, column=3, value="(직접인건비+제경비)의 20~40% (하한 20% 적용)")
ws5.cell(row=row, column=8, value=f"=ROUND((H{di_row}+H{je_row})*0.2,0)")
box(ws5, row, 1, 11)
gi_row = row
row += 1
ws5.cell(row=row, column=1, value="엔지니어링 총계 (부가세 별도)").font = BF
ws5.cell(row=row, column=8, value=f"=H{di_row}+H{je_row}+H{gi_row}").font = BF
box(ws5, row, 1, 11, fill=SUB_FILL, font=BF)
ENG_TOTAL = f"'엔지니어링 내역'!H{row}"

ws5.cell(row=WAGE_ROW_IT - 1, column=1, value="SW기술자 평균임금 공표 (2025)").font = BF
ws5.cell(row=WAGE_ROW_IT - 1, column=3, value="월 20.5일 기준 · 1일 = 0.048 M/M").font = SMALL
ws5.cell(row=WAGE_ROW_IT, column=1, value="IT지원기술자")
ws5.cell(row=WAGE_ROW_IT, column=2, value=245535)
ws5.cell(row=WAGE_ROW_DEV, column=1, value="응용SW개발자")
ws5.cell(row=WAGE_ROW_DEV, column=2, value=337061)
ws5.cell(row=WAGE_ROW_OPS, column=1, value="정보시스템운용자")
ws5.cell(row=WAGE_ROW_OPS, column=2, value=492943)
for rr in range(WAGE_ROW_IT, WAGE_ROW_OPS + 1):
    box(ws5, rr, 1, 2)
widths(ws5, [14, 24, 52, 9, 11, 13, 13, 14, 14, 8, 16])

# ═════════════════════════════════════════════════════════════════
# Sheet 0: 총괄표
# ═════════════════════════════════════════════════════════════════
ws0 = wb.create_sheet("총괄표", 0)
ws0.sheet_properties.tabColor = "FFC000"
ws0.cell(row=1, column=2, value="SLM 스마트 물관리 시스템 — 개발대가 산정 총괄표").font = TITLE_FONT
ws0.cell(row=2, column=2, value="산정 기준: SW사업 대가산정 가이드(KOSA) 단일 FP유형(간이법) + 엔지니어링(SW기술자 평균임금 2025)").font = SMALL
ws0.cell(row=3, column=2, value="기능 기준일: 2026-10-15 (메뉴 57종 · AI 채팅 인텐트 81종 · Migration 0143)").font = SMALL

for c, h in enumerate(["", "구분", "산정 방법", "금액(원)", "비고"], 1):
    ws0.cell(row=5, column=c, value=h)
hdr(ws0, 5, 2, 5)
ws0.cell(row=6, column=2, value="① 소프트웨어 개발비")
ws0.cell(row=6, column=3, value="기능점수(FP) × 단가 × 보정계수 + 이윤")
ws0.cell(row=6, column=4, value="='SW개발비 산정'!I11")
ws0.cell(row=6, column=5, value="FP산정(간이법)·SW개발비 산정 시트")
box(ws0, 6, 2, 5)
ws0.cell(row=7, column=2, value="② 구축·설치 엔지니어링")
ws0.cell(row=7, column=3, value="M/M × SW기술자 평균임금 + 제경비 + 기술료")
ws0.cell(row=7, column=4, value=f"={ENG_TOTAL}")
ws0.cell(row=7, column=5, value="엔지니어링 내역 시트")
box(ws0, 7, 2, 5)
ws0.cell(row=8, column=2, value="합계 (부가세 별도)").font = BF
ws0.cell(row=8, column=4, value="=SUM(D6:D7)").font = BF
box(ws0, 8, 2, 5, fill=TOTAL_FILL, font=BF)
ws0.cell(row=10, column=2, value="※ 기능점수당 단가는 KOSA 2024 가이드(605,784원) 기준 — 최신 공표 단가 확인 후 'SW개발비 산정'!C7 갱신").font = SMALL
ws0.cell(row=11, column=2, value="※ 하드웨어·상용SW 구입비, 유지관리(운영)비는 본 산정 범위 외").font = SMALL
for c in range(2, 6):
    ws0.column_dimensions[get_column_letter(c)].width = [28, 44, 20, 34][c - 2]

# 인쇄 설정
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUTPUT)

# 요약 출력 (수식 미평가 — 파이썬 측 재계산)
cnt = {t: 0 for t in FP_W}
for kind, a, b, c_, d in F:
    if kind != "SEC" and c_ in cnt:
        cnt[c_] += 1
total_fp = sum(cnt[t] * FP_W[t] for t in cnt)
print(f"저장: {OUTPUT}")
print("FP 집계:", {t: cnt[t] for t in cnt}, "| 총", sum(cnt.values()), "개 기능")
print(f"총 FP: {total_fp:.1f}")
size = 1.28 if total_fp < 500 else (1.153 if total_fp > 3000 else 0.4057 * (__import__('math').log(total_fp) - 7.1978) ** 2 + 0.8878)
coef = size * 0.88 * 1.0 * 1.06 * 1.0
dev = total_fp * FP_UNIT_PRICE * coef
print(f"보정계수: 규모 {size:.4f} × 연계 0.88 × 성능 1.0 × 운영환경 1.06 × 보안 1.0 = {coef:.4f}")
print(f"보정 후 개발원가: {dev:,.0f}원 · 이윤 20%: {dev*0.2:,.0f}원 · SW개발비: {dev*1.2:,.0f}원")
eng_di = sum((e[3] * 245535 + e[4] * 337061 + e[5] * 492943) * 0.048 * 20.5 for e in ENG)
eng_total = eng_di + eng_di * 1.2 + (eng_di + eng_di * 1.2) * 0.2
print(f"엔지니어링 직접인건비: {eng_di:,.0f}원 · 총계(제경비 120%+기술료 20%): {eng_total:,.0f}원")
print(f"총괄 합계(부가세 별도): {dev*1.2 + eng_total:,.0f}원")
